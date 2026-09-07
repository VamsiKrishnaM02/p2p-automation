"""
report_builder.py -- flatten processed invoice results into a single Excel
report, one row per invoice (not per source PDF -- a single PDF can split
into multiple invoices, so invoice-level is the correct granularity).

Two entry points depending on where the results came from:
  - flatten_folder_results(all_folder_results)  -- app.py's batch UI
    structure: list of folder_result dicts, each with pdf_results ->
    invoices.
  - flatten_single_pdf_result(result)            -- main.py's own
    process_invoice_pdf() return dict, for a single-PDF CLI run.

Both produce the same flat row shape, then build_excel_report(rows)
writes a formatted .xlsx (bold header, frozen header row, autosized
columns) and returns it as bytes -- ready for st.download_button or
writing straight to disk.
"""

import io
from pathlib import Path
from typing import List, Optional

import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Row shape: one invoice -> one flat dict
# ---------------------------------------------------------------------------
def _invoice_to_row(folder_name: str, source_file: str, inv: dict) -> dict:
    sm = inv.get("sensitive_match") or {}
    lm = inv.get("le_match") or {}
    review_reasons = inv.get("review_reasons") or []
    reasons_str = "; ".join(
        r.get("message", str(r)) if isinstance(r, dict) else str(r)
        for r in review_reasons
    )

    return {
        "Folder": folder_name,
        "Source PDF": source_file,
        "Invoice #": inv.get("invoice_number") or "",
        "Supplier": inv.get("supplier_name") or "",
        "Bill To": inv.get("bill_to_client_name") or "",
        "PO Number(s)": ", ".join(inv.get("po_numbers") or []),
        "Is Sensitive": bool(inv.get("is_sensitive", False)),
        "Sensitive Match": sm.get("matched_name") or "",
        "Sensitive Score": sm.get("score", ""),
        "Sensitive Matched Via Translation": sm.get("translated", False),
        "LE": inv.get("le") or "",
        "Source Site": inv.get("source_site") or "",
        "LE Country": inv.get("le_country") or "",
        "LE Matched Name": lm.get("matched_le_name") or "",
        "LE Score": lm.get("score", ""),
        "LE Matched Via Translation": lm.get("translated", False), 
        "Invoice Type": inv.get("invoice_type") or "",
        "Kofax Email": inv.get("kofax_email") or "",
        "Vietnam Supplier ID": inv.get("vietnam_supplier_id") or "",
        "Vietnam Supplier Name": inv.get("vietnam_supplier_name") or "",
        "Vietnam Renamed": bool(inv.get("vietnam_renamed", False)),
        "Needs Review": bool(inv.get("needs_manual_review", False)),
        "Review Reasons": reasons_str,
        "Output File": Path(inv.get("output_pdf") or "").name,
    }


# ---------------------------------------------------------------------------
# Flattening entry points
# ---------------------------------------------------------------------------
def flatten_folder_results(all_folder_results: List[dict]) -> List[dict]:
    """For app.py's batch UI structure: list of folder_result dicts, each
    with "pdf_results" -> each with "invoices". Skips PDFs that errored
    (they have no invoices to report anyway)."""
    rows = []
    for folder_result in all_folder_results:
        folder_name = folder_result.get("name", "")
        for pr in folder_result.get("pdf_results", []):
            if pr.get("error"):
                continue
            source_file = Path(pr.get("source_file", "")).name
            for inv in pr.get("invoices", []):
                rows.append(_invoice_to_row(folder_name, source_file, inv))
    return rows


def flatten_single_pdf_result(result: dict, folder_name: str = "") -> List[dict]:
    """For main.py's own process_invoice_pdf() return dict -- useful for a
    single-PDF CLI run (see main.py's `if __name__ == "__main__":` block)."""
    source_file = Path(result.get("source_file", "")).name
    return [_invoice_to_row(folder_name, source_file, inv)
            for inv in result.get("invoices", [])]


# ---------------------------------------------------------------------------
# Excel writer
# ---------------------------------------------------------------------------
def build_excel_report(rows: List[dict]) -> bytes:
    """Build a formatted .xlsx (bold header, frozen header row, autosized
    columns) from a flat list of per-invoice row dicts. Returns the file
    as bytes (ready for st.download_button or Path.write_bytes)."""
    df = pd.DataFrame(rows)
    buf = io.BytesIO()

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Invoices")
        ws = writer.sheets["Invoices"]

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"

        for i, col in enumerate(df.columns, start=1):
            max_len = max(
                [len(str(col))] + [len(str(v)) for v in df[col].astype(str)]
            ) if len(df) else len(str(col))
            ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 60)

    buf.seek(0)
    return buf.getvalue()


def save_excel_report(rows: List[dict], out_path: str) -> str:
    """Convenience wrapper: build the report and write it straight to
    disk. Returns the path written."""
    data = build_excel_report(rows)
    Path(out_path).write_bytes(data)
    return out_path